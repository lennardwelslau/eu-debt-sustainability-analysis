# ========================================================================================= #
#               European Commission Debt Sustainability Analysis - Prior Guidance Sheets     #
# ========================================================================================= #
#
# Downloads and parses the European Commission's prior guidance calculation sheets, published
# for each member state on the "Fiscal surveillance in <country>" pages. The sheets contain the
# Commission's own DSA inputs (market forward rates, inflation expectations, debt structure,
# semi-elasticities, stock-flow and ageing cost paths) and results (no-fiscal-policy-change
# baseline, adjustment scenario, required adjustment).
#
# They are used in two ways:
#   1. As default source for inputs that are not publicly available via API (Bloomberg forward
#      rates and inflation swaps, country-specific stock-flow paths), see build.py.
#   2. In 'commission' mode, as the only source for all inputs, to replicate and validate the
#      model against the Commission's results.
#
# Author: Lennard Welslau
# ========================================================================================= #

import re
import time
import hashlib
import datetime
import urllib.parse
import warnings
import numpy as np
import pandas as pd
import requests
import openpyxl

from .schema import RAW_DIR, EU27, COUNTRY_NAMES, SERIES

PRIOR_GUIDANCE_DIR = RAW_DIR / 'commission_prior_guidance'
PAGE_URL = ('https://economy-finance.ec.europa.eu/economic-surveillance-eu-member-states/'
            'country-pages-including-country-reports/{slug}/fiscal-surveillance-{slug}_en')
BASE_URL = 'https://economy-finance.ec.europa.eu'
HEADERS = {'User-Agent': 'Mozilla/5.0 (DSA data pipeline)'}

# Commission two-letter codes used in the sheets
EC_ISO2 = {
    'AUT': 'AT', 'BEL': 'BE', 'BGR': 'BG', 'HRV': 'HR', 'CYP': 'CY', 'CZE': 'CZ', 'DNK': 'DK', 'EST': 'EE',
    'FIN': 'FI', 'FRA': 'FR', 'DEU': 'DE', 'GRC': 'EL', 'HUN': 'HU', 'IRL': 'IE', 'ITA': 'IT', 'LVA': 'LV',
    'LTU': 'LT', 'LUX': 'LU', 'MLT': 'MT', 'NLD': 'NL', 'POL': 'PL', 'PRT': 'PT', 'ROU': 'RO', 'SVK': 'SK',
    'SVN': 'SI', 'ESP': 'ES', 'SWE': 'SE',
}
ISO3_FROM_EC = {v: k for k, v in EC_ISO2.items()}


# ========================================================================================= #
#                                          DOWNLOAD                                         #
# ========================================================================================= #

def _get(url, retries=3, timeout=90):
    for attempt in range(retries):
        try:
            r = requests.get(url, headers=HEADERS, timeout=timeout)
            r.raise_for_status()
            return r
        except requests.RequestException:
            if attempt == retries - 1:
                raise
            time.sleep(5)


def find_prior_guidance_links(country):
    """
    Return download links of all prior guidance calculation sheets on the country's fiscal surveillance page.
    """
    slug = COUNTRY_NAMES[country].lower()
    html = _get(PAGE_URL.format(slug=slug)).text
    links = set()
    for href in re.findall(r'href="([^"]+)"', html):
        href = href.replace('&amp;', '&')
        name = urllib.parse.unquote(href.split('filename=')[-1].split('&')[0]) if 'filename=' in href else ''
        if re.search(r'prior[ _-]guidance', name, re.I) and re.search(r'\.xls[xm]?$', name, re.I):
            links.add(BASE_URL + href if href.startswith('/') else href)
    return sorted(links)


def download_prior_guidance(countries=EU27, dest=PRIOR_GUIDANCE_DIR, overwrite=False, verbose=True):
    """
    Download all prior guidance calculation sheets for the given countries.
    Files are saved as <ISO3>__<original file name> and listed in manifest.csv (url, date, sha256).
    """
    dest.mkdir(parents=True, exist_ok=True)
    manifest_path = dest / 'manifest.csv'
    manifest = pd.read_csv(manifest_path) if manifest_path.exists() else pd.DataFrame(
        columns=['country', 'file', 'url', 'downloaded', 'sha256'])
    for country in countries:
        try:
            links = find_prior_guidance_links(country)
        except Exception as e:
            warnings.warn(f'{country}: could not read fiscal surveillance page ({e})')
            continue
        if not links:
            warnings.warn(f'{country}: no prior guidance calculation sheet found')
        for url in links:
            name = urllib.parse.unquote(url.split('filename=')[-1].split('&')[0])
            name = re.sub(r'[^\w\-. ]', '_', name)
            path = dest / f'{country}__{name}'
            if path.exists() and not overwrite:
                continue
            content = _get(url).content
            path.write_bytes(content)
            manifest = manifest.loc[manifest['file'] != path.name]
            manifest.loc[len(manifest)] = [country, path.name, url, datetime.date.today().isoformat(),
                                           hashlib.sha256(content).hexdigest()]
            if verbose:
                print(f'{country}: downloaded {path.name}')
    manifest.sort_values(['country', 'file']).to_csv(manifest_path, index=False)
    return manifest


# ========================================================================================= #
#                                           PARSE                                           #
# ========================================================================================= #

def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else np.nan


def _read_rows(ws):
    """
    Read a scenario/input sheet into an ordered list of (label, {year: value}, first value column C).
    Years are taken from the first header row with at least ten consecutive years.
    """
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    year_cols = None
    for r in grid:
        cols = {j: int(v) for j, v in enumerate(r) if isinstance(v, (int, float)) and 1990 < v < 2100}
        if len(cols) >= 10:
            year_cols = cols
            break
    rows = []
    for r in grid:
        label = r[1] if len(r) > 1 else None
        if not isinstance(label, str) or not label.strip():
            continue
        values = {y: _num(r[j]) for j, y in year_cols.items() if j < len(r)}
        rows.append((' '.join(label.split()), values, r[2] if len(r) > 2 else None))
    return rows, sorted(year_cols.values())


class PriorGuidanceSheet:
    """
    Parsed prior guidance calculation sheet. Rows are accessed by label pattern and occurrence.
    """

    def __init__(self, path):
        self.path = path
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        self.sheets = {}
        for name in wb.sheetnames:
            if name in ['Input data', 'Baseline NFPC', 'Adjustment scenario', 'Adjust. no safeguard']:
                self.sheets[name] = _read_rows(wb[name])
        self.criteria = [list(r) for r in wb['Criteria results'].iter_rows(values_only=True)]
        wb.close()

        rows, self.years = self.sheets['Input data']
        self.iso2 = str(self._scalar('Input data', r'^Country$')).strip()
        self.country = ISO3_FROM_EC.get(self.iso2, self.iso2)
        self.T = int(self._scalar('Input data', r'^Last year before the adjustment'))
        self.adjustment_end = int(self._scalar('Input data', r'^Last year of adjustment'))
        forecast = next((l for l, _, _ in rows if l.startswith('Key fiscal variables')), '')
        m = re.search(r'\("?([AS]F \d{4})"?\)', forecast)
        self.forecast = m.group(1) if m else forecast
        # Spring forecasts cover the current and next year, autumn forecasts one year more
        year = int(self.forecast[-4:]) if m else self.T
        self.last_forecast_year = year + (2 if self.forecast.startswith('AF') else 1)

    def __repr__(self):
        return f'PriorGuidanceSheet({self.country}, T={self.T}, forecast={self.forecast})'

    def _find(self, sheet, pattern, occurrence=0):
        hits = [(v, c) for l, v, c in self.sheets[sheet][0] if re.search(pattern, l, re.I)]
        if len(hits) <= occurrence:
            raise KeyError(f"'{pattern}' (occurrence {occurrence}) not found in sheet '{sheet}' of {self.path}")
        return hits[occurrence]

    def row(self, sheet, pattern, occurrence=0):
        """Return a row as pd.Series indexed by year."""
        return pd.Series(self._find(sheet, pattern, occurrence)[0], dtype=float)

    def _scalar(self, sheet, pattern, occurrence=0):
        return self._find(sheet, pattern, occurrence)[1]

    def scalar(self, pattern, occurrence=0):
        """Return a scalar from column C of the Input data sheet."""
        return _num(self._scalar('Input data', pattern, occurrence))

    def criteria_value(self, pattern, col, occurrence=0):
        """Return a value from the Criteria results sheet: row by label in column D, value by column letter."""
        j = openpyxl.utils.column_index_from_string(col) - 1
        hits = [r for r in self.criteria if len(r) > 3 and isinstance(r[3], str) and re.search(pattern, r[3], re.I)]
        return _num(hits[occurrence][j]) if len(hits) > occurrence else np.nan

    def criteria_numbers(self, pattern, occurrence=-1):
        """Return the numeric values to the right of a label anywhere in the Criteria results sheet."""
        hits = []
        for r in self.criteria:
            for j, v in enumerate(r[:20]):  # tables span columns A to T, helper values sit in column V
                if isinstance(v, str) and re.search(pattern, ' '.join(v.split()), re.I):
                    hits.append([_num(x) for x in r[j + 1:20] if not np.isnan(_num(x))])
                    break
        return hits[occurrence] if hits else []

    def results(self):
        """
        Commission results used for validation. The SPB at the end of the adjustment (4 and 7 years) is
        reported directly for countries receiving technical information, and derived from the reference
        trajectory (annual change in SPB) otherwise.
        """
        spb_T = self.row('Input data', r'^Structural primary balance').get(self.T, np.nan)
        # Reference trajectory (Table 6): average over the plan period followed by the annual steps
        ref_4y = self.criteria_numbers(r'^Annual change in SPB$', 0)
        ref_7y = self.criteria_numbers(r'^Annual change in SPB$', 1)
        has_technical = any(isinstance(v, str) and 'Technical information' in v for r in self.criteria for v in r)
        technical = self.criteria_numbers(r'^SPB at the end of the adjustment$') if has_technical else []
        if len(ref_4y) >= 5 and len(ref_7y) >= 8:
            kind, spb_end_4y, spb_end_7y = 'reference trajectory', spb_T + sum(ref_4y[1:5]), spb_T + sum(ref_7y[1:8])
        elif len(technical) >= 2:
            kind, spb_end_4y, spb_end_7y = 'technical information', technical[0], technical[1]
        else:
            kind, spb_end_4y, spb_end_7y = 'not found', np.nan, np.nan
        return {
            'country': self.country,
            'T': self.T,
            'forecast': self.forecast,
            'guidance': kind,
            'spb_T': spb_T,
            'spb_end_4y': spb_end_4y,
            'spb_end_7y': spb_end_7y,
            'annual_adjustment_dsa_4y': self.criteria_value(r'^Annual adjustment$', 'G'),
            'annual_adjustment_dsa_7y': self.criteria_value(r'^Annual adjustment$', 'H'),
            'debt_nfpc': self.row('Baseline NFPC', r'^Gross debt'),
            'debt_adjustment': self.row('Adjustment scenario', r'^Gross debt'),
            'spb_adjustment': self.row('Adjustment scenario', r'^Structural primary balance'),
        }


def load_prior_guidance(countries=EU27, vintage='2024', directory=PRIOR_GUIDANCE_DIR):
    """
    Parse the downloaded prior guidance sheets and select one per country.

    vintage='2024':   the first guidance issued for the 2024 fiscal-structural plans (earliest T and forecast)
    vintage='latest': the most recent guidance available for each country
    """
    selected = {}
    for country in countries:
        files = sorted(directory.glob(f'{country}__*.xls*'))
        if not files:
            raise FileNotFoundError(f'No prior guidance sheet for {country} in {directory}. Run download_prior_guidance().')
        sheets = [PriorGuidanceSheet(f) for f in files]
        # Order by reference year and forecast (spring before autumn), then prefer non-updated files
        key = lambda s: (s.T, s.forecast[-4:], s.forecast[:2] == 'AF', 'updated' in s.path.name.lower())
        sheets.sort(key=key)
        selected[country] = sheets[0] if vintage == '2024' else sheets[-1]
    return selected


# ========================================================================================= #
#                                  MAPPING TO MODEL INPUTS                                  #
# ========================================================================================= #

def _prov_frame(series, value):
    return pd.DataFrame(np.where(series.notna(), value, None), index=series.index, columns=series.columns)


def commission_parameters(sheet):
    """
    Parameters from a prior guidance sheet. Returns (params, sources).
    """
    s = sheet
    src = f'Commission prior guidance ({s.forecast}, T={s.T})'
    rows = {
        'FISCAL_MULTIPLIER': s.scalar(r'^Fiscal multiplier'),
        'BUDGET_BALANCE_ELASTICITY': s.scalar(r'^Budget balance semi-elasticity'),
        'DEBT_ST_SHARE': s.scalar(r'^Share of short-term debt in total government debt'),
        'DEBT_LT_MATURING_SHARE': s.row('Baseline NFPC', r'^Share of long-term debt that matures every year').get(s.T, np.nan),
        'DEBT_LT_MATURING_AVG_SHARE': s.scalar(r'^Share of long-term debt that matures every year'),
        'INTEREST_RATE_LT_T10': s.scalar(r'^Long-term nominal interest rate \(T\+10'),
        'INTEREST_RATE_ST_T10': s.scalar(r'^Short-term nominal interest rate \(T\+10'),
        'INTEREST_RATE_LT_T30': s.scalar(r'^Long-term nominal interest rate \(T\+30'),
        'INTEREST_RATE_ST_T30': s.scalar(r'^Short-term nominal interest rate \(T\+30'),
        'INFLATION_T10': s.scalar(r'^GDP deflator .*\(T\+10'),
        'INFLATION_T30': s.scalar(r'^GDP deflator .*\(T\+30'),
        'EXCESSIVE_DEFICIT_PROCEDURE': s.row('Input data', r'^Subject to an excessive deficit procedure').get(s.T, np.nan),
        'LAST_FORECAST_YEAR': s.last_forecast_year,
        'FISCAL_MULTIPLIER_PERSISTENT': 0,
    }
    sources = {k: ('commission', src) for k in rows}
    sources['INTEREST_RATE_LT_T10'] = ('commission', f'{src}: Bloomberg country-specific forward rates')
    sources['INTEREST_RATE_ST_T10'] = ('commission', f'{src}: Bloomberg forward rates')
    sources['INFLATION_T10'] = ('commission', f'{src}: Bloomberg inflation swaps')
    return rows, sources


def commission_stock_flow(sheet):
    """
    Stock-flow adjustment path (% of GDP, total incl. exchange rate effects) from a prior guidance sheet.
    """
    return sheet.row('Input data', r'^Stock-flow adjustment \(total\)')


def build_country_commission(sheet, end_year=2070):
    """
    Map a prior guidance sheet to model inputs (commission mode). All inputs come from the sheet.
    Paths beyond the sheet horizon (T+17) are held constant; levels are scaled to real GDP in T.
    """
    s = sheet
    T = s.T
    src = f'Commission prior guidance ({s.forecast}, T={T})'
    years = range(min(s.years), end_year + 1)
    out = pd.DataFrame(index=pd.Index(years, name='YEAR'), columns=list(SERIES), dtype=float)
    upto = lambda row, last: row.loc[:last]
    nfpc = 'Baseline NFPC'

    # Growth and inflation paths (full horizon of the sheet)
    out['REAL_GDP'] = s.row(nfpc, r'^Level$', 0)
    out['REAL_GDP_GROWTH'] = s.row(nfpc, r'^Growth rate$', 0)
    out['POTENTIAL_GDP'] = s.row(nfpc, r'^Level$', 1)
    out['POTENTIAL_GDP_GROWTH'] = s.row('Input data', r'^Growth rate$', 1)
    out['GDP_DEFLATOR_PCH'] = s.row(nfpc, r'^GDP deflator')
    out['NOMINAL_GDP_GROWTH'] = upto(s.row(nfpc, r'^Growth rate$', 1), T + 2)

    # Nominal levels are not reported: index nominal GDP to real GDP in T, only ratios matter for the DSA
    ng = s.row(nfpc, r'^Growth rate$', 1)
    ngdp = pd.Series(np.nan, index=ng.index)
    ngdp[T] = out.loc[T, 'REAL_GDP']
    for y in range(T + 1, T + 3):
        ngdp[y] = ngdp[y - 1] * (1 + ng[y] / 100)
    out['NOMINAL_GDP'] = ngdp.loc[T:T + 2]

    # Fiscal variables up to T+2 (the model projects from T+1 onwards)
    out['STRUCTURAL_PRIMARY_BALANCE'] = upto(s.row(nfpc, r'^Structural primary balance'), T + 2)
    out['PRIMARY_BALANCE'] = upto(s.row(nfpc, r'^\(1\) Primary balance'), T + 2)
    out['FISCAL_BALANCE'] = upto(s.row(nfpc, r'^Headline balance'), T + 2)
    # Implicit interest rate: forecast up to the last forecast year of the sheet, projected by the model thereafter
    out['IMPLICIT_INTEREST_RATE'] = upto(s.row(nfpc, r'^Nominal implicit interest rate on debt'), s.last_forecast_year)
    # Commission adjustment of the Excel approximation to its Stata model (non-zero for Greece only)
    iir_adj = s.row('Adjust. no safeguard', r'^Diff\. STATA').loc[T + 2:]
    if iir_adj.abs().max() > 1e-6:
        out['IMPLICIT_INTEREST_RATE_ADJ'] = iir_adj
    out.loc[T, 'PRIMARY_EXPENDITURE_SHARE'] = s.scalar(r'^Share of primary expenditure in GDP')
    out['DEBT_RATIO'] = upto(s.row(nfpc, r'^Gross debt'), T + 2)
    out['ONE_OFF_MEASURES'] = s.row('Input data', r'^One-off and other temporary measures')
    out['DEBT_TOTAL'] = out['DEBT_RATIO'] * out['NOMINAL_GDP'] / 100

    # Stock-flow (total, incl. exchange rate effect) as exogenous path; FX effects are captured here
    sf = commission_stock_flow(s)
    out['STOCK_FLOW_RATIO'] = sf
    out['STOCK_FLOW'] = sf * out['NOMINAL_GDP'] / 100
    out.loc[T - 3:T + 2, 'EXR_EUR'] = 1.0
    out.loc[T - 3:T + 2, 'EXR_USD'] = 1.0

    # Market rates: Commission path from T-1 to T+17
    out['INTEREST_RATE_LT'] = s.row(nfpc, r'^Long-term interest rate')
    out['INTEREST_RATE_ST'] = s.row(nfpc, r'^Short-term interest rate')
    out['DEBT_LT_MATURING_SHARE_PATH'] = s.row(nfpc, r'^Share of long-term debt that matures every year')

    # Ageing costs net of pension taxes and property income; levels held constant after T+17
    out['AGEING_COST'] = s.row('Input data', r'^Total ageing cost')
    out['TAX_AND_PROPERTY_INCOME'] = s.row('Input data', r'^Property income')
    out.loc[out.index < T, 'AGEING_COST'] = np.nan
    last = max(s.years)
    for code in ['POTENTIAL_GDP_GROWTH', 'AGEING_COST', 'TAX_AND_PROPERTY_INCOME']:
        out.loc[last + 1:, code] = out.loc[last, code]

    prov = _prov_frame(out, 'commission')
    for code in ['POTENTIAL_GDP_GROWTH', 'AGEING_COST', 'TAX_AND_PROPERTY_INCOME']:
        prov.loc[last + 1:, code] = 'derived'
    for code in ['NOMINAL_GDP', 'DEBT_TOTAL', 'STOCK_FLOW', 'EXR_EUR', 'EXR_USD']:
        prov.loc[out[code].notna(), code] = 'derived'

    ssrc = {c: src for c in out.columns if out[c].notna().any()}
    ssrc.update({
        'NOMINAL_GDP': 'Index: real GDP in T grown with Commission nominal growth (levels not reported; only ratios matter)',
        'DEBT_TOTAL': 'Debt ratio x indexed nominal GDP',
        'STOCK_FLOW': 'Stock-flow ratio x indexed nominal GDP',
        'EXR_EUR': 'Constant: exchange rate effects are included in the Commission stock-flow adjustment',
        'EXR_USD': 'Constant: exchange rate effects are included in the Commission stock-flow adjustment',
        'AGEING_COST': f'{src}: total ageing cost net of taxes on pensions, constant after {last}',
        'TAX_AND_PROPERTY_INCOME': f'{src}: property income, constant after {last}',
        'POTENTIAL_GDP_GROWTH': f'{src}, constant after {last}',
        'STOCK_FLOW_RATIO': f'{src}: total stock-flow adjustment incl. exchange rate effects',
        'IMPLICIT_INTEREST_RATE': f'{src}: forecast years; projected by the model thereafter',
        'IMPLICIT_INTEREST_RATE_ADJ': f'{src}: difference between Commission Stata model and Excel approximation',
    })

    params, psrc = commission_parameters(s)
    params['DEBT_DOMESTIC_SHARE'] = 1.0
    params['DEBT_EUR_SHARE'] = 0.0
    psrc['DEBT_DOMESTIC_SHARE'] = ('assumption', 'All debt treated as domestic: exchange rate effects are in the stock-flow adjustment')
    psrc['DEBT_EUR_SHARE'] = ('assumption', 'All debt treated as domestic: exchange rate effects are in the stock-flow adjustment')

    return {'params': params, 'param_sources': psrc, 'series': out, 'series_provenance': prov, 'series_sources': ssrc,
            'reference_year': T}
