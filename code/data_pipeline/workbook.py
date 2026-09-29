# ========================================================================================= #
#               European Commission Debt Sustainability Analysis - Input Workbook           #
# ========================================================================================= #
#
# Writes and reads the Excel input workbook used by the DSA model. The workbook contains:
#
#   README     - description, vintage and colour legend
#   Overview   - key starting values for all countries (for quick plausibility checks)
#   Overrides  - optional user overrides (COUNTRY, CODE, YEAR, VALUE, COMMENT), applied on read
#   Sources    - source note for every variable
#   <ISO3>     - one sheet per country with a PARAMETERS block and a TIME SERIES block
#   Shocks_Q / Shocks_A   - historical shocks (first differences) for the stochastic model
#   History_Q / History_A - underlying historical levels (for transparency, not read by the model)
#
# Values can be edited directly in the country sheets or, preferably, entered in the Overrides
# sheet, which keeps user inputs separate from the fetched data and survives a rebuild.
#
# Author: Lennard Welslau
# ========================================================================================= #

import datetime
import warnings
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .schema import PARAMETERS, SERIES, PROVENANCE, OVERRIDE_COLUMNS, COUNTRY_NAMES, resolve_input_path

PARAM_MARKER = 'PARAMETERS'
SERIES_MARKER = 'TIME SERIES'
SERIES_META_COLUMNS = ['Code', 'Group', 'Description', 'Unit', 'Source']
PARAM_COLUMNS = ['Code', 'Description', 'Unit', 'Value', 'Source']

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_THIN = Border(bottom=Side(style='thin', color='BFBFBF'))


def _fill(provenance):
    colour = PROVENANCE.get(provenance, PROVENANCE['derived'])[1]
    return PatternFill('solid', fgColor=colour)


def _header(ws, row, values):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(horizontal='center')


def _clean_value(v):
    if v is None:
        return None
    if isinstance(v, (float, np.floating)) and np.isnan(v):
        return None
    if isinstance(v, np.generic):
        return v.item()
    return v


# ========================================================================================= #
#                                           WRITER                                          #
# ========================================================================================= #

SHOCK_SHEETS = ['Shocks_Q', 'Shocks_A', 'History_Q', 'History_A']


def write_workbook(path, countries, meta, overrides=None, shocks=None):
    """
    Write the input workbook.

    Parameters:
        path (str/Path): Output path.
        countries (dict): {iso3: {
                'params': {code: value},
                'param_sources': {code: (provenance, note)},
                'series': DataFrame (index=year, columns=codes),
                'series_provenance': DataFrame of provenance codes (same shape, optional),
                'series_sources': {code: note}}}
        meta (dict): Workbook metadata, e.g. vintage, mode, description, reference_year.
        overrides (DataFrame): Optional overrides table with OVERRIDE_COLUMNS.
        shocks (dict): Optional {sheet name: long DataFrame (COUNTRY, PERIOD, variables)} for the stochastic model.
    """
    no_history = [iso for iso, d in countries.items() if d.get('reference_year') is not None
                  and pd.isna(d['series']['FISCAL_BALANCE'].get(d['reference_year'] - 1, np.nan))]
    if no_history:
        raise ValueError(f'No fiscal balance in T-1 for {no_history}. The model needs historical data from T-1 '
                         f'(EDP abrogation rule).')

    wb = Workbook()
    _write_readme(wb.active, meta)
    _write_overview(wb.create_sheet('Overview'), countries, meta)
    _write_overrides(wb.create_sheet('Overrides'), overrides)
    _write_sources(wb.create_sheet('Sources'), countries, meta)
    for iso, data in countries.items():
        _write_country(wb.create_sheet(iso), iso, data, meta)
    for name in SHOCK_SHEETS:
        if shocks is not None and name in shocks:
            _write_table(wb.create_sheet(name), shocks[name])
    wb.save(path)
    return path


def _write_readme(ws, meta):
    ws.title = 'README'
    ws['A1'] = 'DSA model input data'
    ws['A1'].font = _TITLE
    rows = [
        ('Vintage', meta.get('vintage')),
        ('Created', meta.get('created', datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))),
        ('Build mode', meta.get('mode')),
        ('Reference year (T)', meta.get('reference_year')),
        ('Description', meta.get('description')),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = _BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='How to use').font = _BOLD
    notes = [
        'Each country sheet holds a PARAMETERS block (one value per country) and a TIME SERIES block (one value per year).',
        'Blank cells are filled by the model following the Commission methodology (interpolation towards T+10 and T+30 anchors).',
        'To change an input, add a row to the Overrides sheet (YEAR blank for parameters). Overrides are applied when the model reads the file.',
        'Values can also be edited directly in the country sheets; such edits are lost when the workbook is rebuilt.',
        'Do not rename the sheets, the block markers in column A, or the header rows.',
    ]
    for n in notes:
        r += 1
        ws.cell(row=r, column=1, value=f'- {n}')
    r += 2
    ws.cell(row=r, column=1, value='Cell colours (data provenance)').font = _BOLD
    for code, (label, _) in PROVENANCE.items():
        r += 1
        ws.cell(row=r, column=1, value=code).fill = _fill(code)
        ws.cell(row=r, column=2, value=label)
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90


def _write_overview(ws, countries, meta):
    T = meta.get('reference_year')
    ws['A1'] = f'Key inputs by country (reference year T = {T})'
    ws['A1'].font = _TITLE
    series_cols = ['DEBT_RATIO', 'STRUCTURAL_PRIMARY_BALANCE', 'FISCAL_BALANCE', 'IMPLICIT_INTEREST_RATE',
                   'REAL_GDP_GROWTH', 'GDP_DEFLATOR_PCH', 'INTEREST_RATE_LT']
    param_cols = ['INTEREST_RATE_LT_T10', 'INFLATION_T10', 'DEBT_ST_SHARE', 'BUDGET_BALANCE_ELASTICITY']
    header = ['Country', 'ISO', 'T'] + [f'{c} (T)' for c in series_cols] + param_cols
    _header(ws, 3, header)
    for i, (iso, data) in enumerate(countries.items(), start=4):
        Tc = meta.get('reference_years', {}).get(iso, T)
        ws.cell(row=i, column=1, value=COUNTRY_NAMES.get(iso, iso))
        ws.cell(row=i, column=2, value=iso)
        ws.cell(row=i, column=3, value=Tc)
        s = data['series']
        for j, c in enumerate(series_cols, start=4):
            v = s.loc[Tc, c] if (Tc in s.index and c in s.columns) else None
            ws.cell(row=i, column=j, value=_clean_value(v)).number_format = '0.00'
        for j, c in enumerate(param_cols, start=4 + len(series_cols)):
            ws.cell(row=i, column=j, value=_clean_value(data['params'].get(c))).number_format = '0.000'
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = 'D4'


def _write_overrides(ws, overrides):
    ws['A1'] = 'User overrides - applied on top of the country sheets when the model reads this file'
    ws['A1'].font = _BOLD
    ws['A2'] = 'Leave YEAR blank for parameters. Example: POL | INTEREST_RATE_LT_T10 | | 5.5 | own assumption'
    _header(ws, 4, OVERRIDE_COLUMNS)
    if overrides is not None:
        for i, row in enumerate(overrides[OVERRIDE_COLUMNS].itertuples(index=False), start=5):
            for j, v in enumerate(row, start=1):
                ws.cell(row=i, column=j, value=_clean_value(v)).fill = _fill('assumption')
    for col, w in zip('ABCDE', [10, 32, 8, 12, 60]):
        ws.column_dimensions[col].width = w


def _write_sources(ws, countries, meta):
    ws['A1'] = 'Sources by variable'
    ws['A1'].font = _TITLE
    _header(ws, 3, ['Code', 'Kind', 'Description', 'Unit', 'Source (first country listed; see country sheets for details)'])
    r = 4
    first = next(iter(countries.values()))
    for code, (desc, unit) in PARAMETERS.items():
        note = first.get('param_sources', {}).get(code, (None, ''))[1]
        for j, v in enumerate([code, 'parameter', desc, unit, note], start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    for code, (desc, unit, _) in SERIES.items():
        note = first.get('series_sources', {}).get(code, '')
        for j, v in enumerate([code, 'series', desc, unit, note], start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    for code, note in meta.get('shock_sources', {}).items():
        for j, v in enumerate([code, 'shocks', 'Historical series for stochastic shocks (Shocks_Q, Shocks_A)', '', note], start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    for col, w in zip('ABCDE', [30, 10, 70, 22, 110]):
        ws.column_dimensions[col].width = w


def _write_table(ws, df):
    _header(ws, 1, list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([_clean_value(v) for v in row])
    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 20 if j > 2 else 10
    ws.freeze_panes = 'C2'


def _write_country(ws, iso, data, meta):
    ws['A1'] = f"{COUNTRY_NAMES.get(iso, iso)} ({iso}) - DSA input data"
    ws['A1'].font = _TITLE
    ws['A2'] = f"Vintage: {meta.get('vintage')} | Mode: {meta.get('mode')} | Reference year T: {meta.get('reference_year')}"

    # Parameters block
    r = 4
    ws.cell(row=r, column=1, value=PARAM_MARKER).font = _BOLD
    r += 1
    _header(ws, r, PARAM_COLUMNS)
    params, psources = dict(data['params']), dict(data.get('param_sources', {}))
    if data.get('reference_year') is not None:
        params['REFERENCE_YEAR'] = data['reference_year']
        psources['REFERENCE_YEAR'] = ('derived', 'Last year before the forecast horizon of the fiscal data')
    for code, (desc, unit) in PARAMETERS.items():
        r += 1
        prov, note = psources.get(code, ('derived', ''))
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=unit)
        c = ws.cell(row=r, column=4, value=_clean_value(params.get(code)))
        c.fill, c.number_format = _fill(prov), ('0' if code == 'REFERENCE_YEAR' else '0.0000')
        ws.cell(row=r, column=5, value=note)

    # Time series block
    r += 2
    ws.cell(row=r, column=1, value=SERIES_MARKER).font = _BOLD
    r += 1
    series = data['series']
    prov_df = data.get('series_provenance')
    ssources = data.get('series_sources', {})
    years = [int(y) for y in series.index]
    _header(ws, r, SERIES_META_COLUMNS + years)
    header_row = r
    for code, (desc, unit, group) in SERIES.items():
        r += 1
        for j, v in enumerate([code, group, desc, unit, ssources.get(code, '')], start=1):
            ws.cell(row=r, column=j, value=v).border = _THIN
        for j, y in enumerate(years, start=len(SERIES_META_COLUMNS) + 1):
            v = series.loc[y, code] if code in series.columns else None
            c = ws.cell(row=r, column=j, value=_clean_value(v))
            c.number_format, c.border = '0.000', _THIN
            if _clean_value(v) is not None:
                prov = prov_df.loc[y, code] if (prov_df is not None and code in prov_df.columns) else 'derived'
                c.fill = _fill(prov)

    # Layout
    for col, w in zip('ABCDE', [30, 10, 60, 20, 45]):
        ws.column_dimensions[col].width = w
    for j in range(len(SERIES_META_COLUMNS) + 1, len(SERIES_META_COLUMNS) + len(years) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 9
    ws.freeze_panes = ws.cell(row=header_row + 1, column=len(SERIES_META_COLUMNS) + 1)


# ========================================================================================= #
#                                           READER                                          #
# ========================================================================================= #

_CACHE = {}


def read_workbook(input_file):
    """
    Read the full input workbook into {iso3: (series DataFrame, params dict)} and a meta dict.
    Results are cached by file path and modification time.
    """
    path = resolve_input_path(input_file)
    key = (str(path), path.stat().st_mtime)
    if key not in _CACHE:
        sheets = pd.read_excel(path, sheet_name=None, header=None)
        countries = {name: _parse_country_sheet(df) for name, df in sheets.items() if name in COUNTRY_NAMES}
        _apply_overrides(countries, sheets.get('Overrides'))
        meta = _parse_readme(sheets.get('README'))
        meta['shocks'] = {name[-1]: _parse_table(sheets[name]) for name in ['Shocks_Q', 'Shocks_A'] if name in sheets}
        _CACHE.clear()
        _CACHE[key] = (countries, meta)
    return _CACHE[key]


def read_shocks(input_file, country, frequency='Q'):
    """
    Return historical shocks (first differences) for one country, indexed by quarterly or annual period.
    """
    _, meta = read_workbook(input_file)
    freq = frequency[0].upper()
    if freq not in meta['shocks']:
        raise KeyError(f'No Shocks_{freq} sheet in {input_file}. Rebuild the workbook with shock data.')
    df = meta['shocks'][freq]
    df = df.loc[df['COUNTRY'] == country].drop(columns='COUNTRY').set_index('PERIOD')
    df.index = pd.PeriodIndex(df.index.astype(str), freq='Q' if freq == 'Q' else 'Y')
    return df.astype(float)


def _parse_table(df):
    df = df.copy()
    df.columns = df.iloc[0]
    return df.iloc[1:].reset_index(drop=True)


def read_country(input_file, country):
    """
    Return (series DataFrame indexed by year, params dict) for one country, with overrides applied.
    Copies are returned so that models cannot modify the cached data.
    """
    countries, _ = read_workbook(input_file)
    if country not in countries:
        raise KeyError(f'No sheet for {country} in {input_file}. Available: {sorted(countries)}')
    series, params = countries[country]
    return series.copy(), dict(params)


def _find_row(df, marker):
    hits = df.index[df.iloc[:, 0].astype(str).str.strip() == marker]
    if len(hits) == 0:
        raise ValueError(f"Marker '{marker}' not found in column A")
    return hits[0]


def _parse_country_sheet(df):
    # Parameters: rows between the parameter header and the series marker (blank rows are skipped)
    params = {}
    series_marker = _find_row(df, SERIES_MARKER)
    for r in range(_find_row(df, PARAM_MARKER) + 2, series_marker):
        if pd.notna(df.iat[r, 0]) and str(df.iat[r, 0]).strip():
            params[str(df.iat[r, 0]).strip()] = pd.to_numeric(df.iat[r, 3], errors='coerce')

    # Series: header row lists meta columns followed by integer years; all rows below (blank rows are skipped)
    h = series_marker + 1
    header = df.iloc[h].tolist()
    year_cols = [(j, int(v)) for j, v in enumerate(header) if isinstance(v, (int, float, np.integer, np.floating))
                 and pd.notna(v) and float(v).is_integer()]
    data = {}
    for r in range(h + 1, len(df)):
        if pd.notna(df.iat[r, 0]) and str(df.iat[r, 0]).strip():
            data[str(df.iat[r, 0]).strip()] = [pd.to_numeric(df.iat[r, j], errors='coerce') for j, _ in year_cols]
    series = pd.DataFrame(data, index=pd.Index([y for _, y in year_cols], name='YEAR'), dtype=float)

    # Rows or parameters missing from the sheet (e.g. optional inputs in older workbooks) are treated as empty
    series = series.reindex(columns=list(dict.fromkeys(list(series.columns) + list(SERIES))))
    for code in PARAMETERS:
        params.setdefault(code, np.nan)
    return series, params


def _check_override(country, code, year, value, countries):
    """
    Validate one override row. Returns (ISO3, code, year or None, value); raises ValueError with the reason.
    Parameters take no YEAR, series require one; VALUE must be numeric.
    """
    iso, code = str(country).strip().upper(), str(code).strip()
    row = f'Override {iso} | {code} | {year} | {value}'
    if iso not in countries:
        raise ValueError(f'{row}: unknown country (workbook has {sorted(countries)})')
    if year is None or (isinstance(year, str) and not year.strip()) or (not isinstance(year, str) and pd.isna(year)):
        year = None
        if code not in PARAMETERS:
            raise ValueError(f'{row}: YEAR is empty, so CODE must be a parameter (see the PARAMETERS block)')
    else:
        try:
            year_float = float(year)
        except (TypeError, ValueError):
            raise ValueError(f'{row}: YEAR must be a year') from None
        if not year_float.is_integer():
            raise ValueError(f'{row}: YEAR must be a whole year')
        year = int(year_float)
        if code not in SERIES:
            raise ValueError(f'{row}: YEAR is given, so CODE must be a time series (see the TIME SERIES block)')
    number = pd.to_numeric(value, errors='coerce')
    if pd.isna(number):
        raise ValueError(f'{row}: VALUE must be a number (use a decimal point)')
    return iso, code, year, float(number)


def _override_table(df):
    """
    Overrides table below the header row of the Overrides sheet (header located by the COUNTRY column label).
    """
    h = df.index[df.iloc[:, 0].astype(str).str.strip() == OVERRIDE_COLUMNS[0]]
    if len(h) == 0:
        return None
    table = df.iloc[h[0] + 1:, :len(OVERRIDE_COLUMNS)].copy()
    table.columns = OVERRIDE_COLUMNS
    return table.dropna(subset=['COUNTRY', 'CODE'])


def _apply_overrides(countries, df):
    table = _override_table(df) if df is not None else None
    if table is None:
        return
    seen = set()
    for row in table.itertuples(index=False):
        iso, code, year, value = _check_override(row.COUNTRY, row.CODE, row.YEAR, row.VALUE, countries)
        if (iso, code, year) in seen:
            warnings.warn(f'Duplicate override for {iso} {code} {year or ""}: the last row is used')
        seen.add((iso, code, year))
        series, params = countries[iso]
        if year is None:
            params[code] = value
        else:
            if year not in series.index:
                series.loc[year] = np.nan
                series.sort_index(inplace=True)
            series.loc[year, code] = value


def apply_overrides(series, params, overrides, country):
    """
    Apply user inputs to the data of one country (in memory, the workbook is not changed).

    Parameters:
        series (DataFrame), params (dict): Country data as returned by read_country (modified in place).
        overrides (DataFrame or list of dicts): Rows with CODE, VALUE and YEAR (omitted or None for parameters);
            COUNTRY is optional (rows for other countries are ignored), COMMENT is ignored.
    """
    table = pd.DataFrame(overrides)
    if 'COUNTRY' not in table:
        table['COUNTRY'] = country
    table = table.reindex(columns=OVERRIDE_COLUMNS)
    table['COUNTRY'] = table['COUNTRY'].fillna(country)
    for r in table.itertuples(index=False):
        if str(r.COUNTRY).strip().upper() != country:
            continue
        _, code, year, value = _check_override(r.COUNTRY, r.CODE, r.YEAR, r.VALUE, [country])
        if year is None:
            params[code] = value
        else:
            if year not in series.index:
                series.loc[year] = np.nan
                series.sort_index(inplace=True)
            series.loc[year, code] = value
    return series, params


def set_overrides(input_file, overrides, out_file=None, replace=False):
    """
    Add user assumptions to the Overrides sheet of an input workbook (all other sheets are kept as they are).

    Parameters:
        input_file (str/Path): Workbook name in data/InputData or path.
        overrides (DataFrame or list of dicts): Rows with COUNTRY, CODE, YEAR (blank/None for parameters), VALUE, COMMENT.
        out_file (str/Path): Save to this file instead of overwriting input_file (name in data/InputData or path).
        replace (bool): If True, existing overrides are removed first.

    Returns:
        Path of the saved workbook.
    """
    from openpyxl import load_workbook
    path = resolve_input_path(input_file)
    out = path if out_file is None else resolve_input_path(out_file, must_exist=False)

    # Validate all rows before writing (countries are the country sheets of the workbook)
    wb = load_workbook(path)
    countries = [name for name in wb.sheetnames if len(name) == 3 and name.isupper()]
    table = pd.DataFrame(overrides).reindex(columns=OVERRIDE_COLUMNS)
    checked = [_check_override(r.COUNTRY, r.CODE, r.YEAR, r.VALUE, countries) for r in table.itertuples(index=False)]
    table[['COUNTRY', 'CODE', 'YEAR', 'VALUE']] = pd.DataFrame(checked, index=table.index).astype(object).values

    # First data row: below the COUNTRY header row
    ws = wb['Overrides']
    header = next((r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == OVERRIDE_COLUMNS[0]), None)
    if header is None:
        raise ValueError(f'No overrides header ({OVERRIDE_COLUMNS[0]}) found in the Overrides sheet of {path}')
    first = header + 1
    if replace:
        ws.delete_rows(first, max(ws.max_row - first + 1, 0))
    row = max(first, ws.max_row + 1)
    while row > first and all(ws.cell(row=row - 1, column=j).value is None for j in range(1, 6)):
        row -= 1
    for values in table.itertuples(index=False):
        for j, v in enumerate(values, start=1):
            ws.cell(row=row, column=j, value=_clean_value(v)).fill = _fill('assumption')
        row += 1
    wb.save(out)
    return out


def _parse_readme(df):
    meta = {}
    if df is None:
        return meta
    for _, row in df.iterrows():
        k, v = row.iloc[0], row.iloc[1] if len(row) > 1 else None
        if isinstance(k, str) and pd.notna(v):
            meta[k.strip()] = v
    return meta
