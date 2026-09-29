# ========================================================================================= #
#               European Commission Debt Sustainability Analysis - Data Sources             #
# ========================================================================================= #
#
# Fetchers for public data used in 'api' mode. Each function returns tidy pandas objects and
# raises on failure, so that missing data is reported rather than silently filled.
#
#   AMECO            via DBnomics API          Commission forecast T to T+2
#   ECB Data Portal  data-api.ecb.europa.eu    benchmark rates, debt structure (GFS)
#   ESM              RepaymentData.csv         ESM/EFSF loan repayment schedules
#   Ageing Report    2024 statistical annex    ageing costs, long-term potential growth
#   DSM              country fiches            stock-flow exceptions, taxes and property income
#   OGWG             user-supplied file        real and potential GDP to T+5 (optional)
#
# Author: Lennard Welslau
# ========================================================================================= #

import io
import re
import time
import numpy as np
import pandas as pd
import requests
import openpyxl

from .schema import RAW_DIR
from .commission import EC_ISO2

HEADERS = {'User-Agent': 'Mozilla/5.0 (DSA data pipeline)'}
REFERENCE_DIR = RAW_DIR / 'reference'

REFERENCE_FILES = {
    'awg': ('AWG2024_horizontal_tables.xlsx',
            'https://economy-finance.ec.europa.eu/document/download/403cc04f-9487-406b-a48a-22538e0d461c_en'
            '?filename=2024_Ageing_Report-Statistical_annex_all_horizontal_tables.xlsx'),
    'dsm': ('DSM2025_country_fiches.xlsx',
            'https://economy-finance.ec.europa.eu/document/download/19852ffb-d8c5-4d47-a402-aa6179dc8051_en'
            '?filename=DSM%202025%20country%20fiches%20tables%20and%20graphs.xlsx'),
}
ESM_URL = 'https://www.esm.europa.eu/financial-assistance/programme-database/reports/RepaymentData.csv'
ECB_URL = 'https://data-api.ecb.europa.eu/service/data/'

# ISO2 codes used by the ECB and ESM (Greece is GR, not EL)
ISO2 = {k: ('GR' if v == 'EL' else v) for k, v in EC_ISO2.items()}

# Changes in EDP status after the Commission prior guidance sheets (whose flag refers to the sheet's reference year):
# {ISO3: (status, first reference year the status applies to, note)}. Applied in API mode if the year is <= T.
EDP_STATUS_UPDATES = {
    'AUT': (1, 2025, 'Excessive deficit procedure opened in 2025, after the Commission 2024 prior guidance'),
}

# Euro adoption year, used to decide whether euro-denominated debt counts as domestic
EURO_ADOPTION = {
    'AUT': 1999, 'BEL': 1999, 'DEU': 1999, 'ESP': 1999, 'FIN': 1999, 'FRA': 1999, 'IRL': 1999, 'ITA': 1999,
    'LUX': 1999, 'NLD': 1999, 'PRT': 1999, 'GRC': 2001, 'SVN': 2007, 'CYP': 2008, 'MLT': 2008, 'SVK': 2009,
    'EST': 2011, 'LVA': 2014, 'LTU': 2015, 'HRV': 2023, 'BGR': 2026,
}
CURRENCY = {'BGR': 'BGN', 'CZE': 'CZK', 'DNK': 'DKK', 'HUN': 'HUF', 'POL': 'PLN', 'ROU': 'RON', 'SWE': 'SEK'}


def _get(url, retries=3, timeout=120, headers=None, **kwargs):
    for attempt in range(retries):
        try:
            r = requests.get(url, headers={**HEADERS, **(headers or {})}, timeout=timeout, **kwargs)
            r.raise_for_status()
            return r
        except requests.RequestException:
            if attempt == retries - 1:
                raise
            time.sleep(5)


def download_reference_files(overwrite=False):
    """
    Download published reference files (Ageing Report annex, DSM country fiches) to data/RawData/reference.
    """
    REFERENCE_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}
    for key, (name, url) in REFERENCE_FILES.items():
        path = REFERENCE_DIR / name
        if overwrite or not path.exists():
            path.write_bytes(_get(url).content)
        paths[key] = path
    return paths


# ========================================================================================= #
#                                     AMECO (via DBnomics)                                  #
# ========================================================================================= #

AMECO = {  # code -> AMECO variable key
    'NOMINAL_GDP': '1.0.0.0.UVGD',
    'REAL_GDP': '1.1.0.0.OVGD',
    'POTENTIAL_GDP': '1.0.0.0.OVGDP',
    'GDP_DEFLATOR': '3.1.0.0.PVGD',
    'PRIMARY_BALANCE': '1.0.319.0.UBLGIE',
    'FISCAL_BALANCE': '1.0.319.0.UBLGE',
    'STRUCTURAL_PRIMARY_BALANCE': '1.0.319.0.UBLGBPS',
    'ONE_OFF_MEASURES': '1.0.319.0.UOOMS',
    'PRIMARY_EXPENDITURE_SHARE': '1.0.319.0.UUTGI',
    'IMPLICIT_INTEREST_RATE': '1.0.0.0.AYIGD',
    'DEBT_TOTAL': '1.0.0.0.UDGG',
    'DEBT_RATIO': '1.0.319.0.UDGG',
    'STOCK_FLOW': '1.0.0.0.UDGGS',
    'XNE': '1.0.99.0.XNE',
}


def fetch_ameco(countries, start_year=2000):
    """
    Fetch AMECO series for all countries from DBnomics.
    Returns (DataFrame with MultiIndex columns (country, code) and year index, vintage string).
    """
    from dbnomics import fetch_series
    area = {c: ('ROM' if c == 'ROU' else c) for c in list(countries) + ['USA', 'EA20']}
    ids = [f'AMECO/{key.split(".")[-1]}/{area[c]}.{key}' for c in area for code, key in AMECO.items()
           if c not in ['USA', 'EA20'] or code in (['XNE'] if c == 'USA' else ['GDP_DEFLATOR'])]
    frames = []
    for i in range(0, len(ids), 50):
        f = fetch_series(ids[i:i + 50])
        f = f.loc[:, ~f.columns.duplicated()]
        frames.append(f[[c for c in ['series_code', 'period', 'value', 'indexed_at'] if c in f.columns]])
    raw = pd.concat(frames, ignore_index=True)
    raw['year'] = pd.to_datetime(raw['period']).dt.year
    raw = raw.loc[raw['year'] >= start_year]
    inv_area = {v: k for k, v in area.items()}
    inv_key = {v: k for k, v in AMECO.items()}
    raw['country'] = raw['series_code'].str.split('.').str[0].map(inv_area)
    raw['code'] = raw['series_code'].str.split('.', n=1).str[1].map(inv_key)
    df = raw.pivot_table(index='year', columns=['country', 'code'], values='value')
    vintage = str(raw['indexed_at'].max())[:10] if 'indexed_at' in raw.columns else 'unknown'
    return df, vintage


# ========================================================================================= #
#                                        ECB Data Portal                                    #
# ========================================================================================= #

def fetch_ecb(flow, key):
    """
    Fetch an ECB series (wildcards and '+' combinations allowed) as tidy DataFrame
    with columns REF_AREA, TIME_PERIOD, OBS_VALUE (+ all key dimensions).
    """
    r = _get(f'{ECB_URL}{flow}/{key}', params={'format': 'csvdata'})
    df = pd.read_csv(io.StringIO(r.text))
    df['OBS_VALUE'] = pd.to_numeric(df['OBS_VALUE'], errors='coerce')
    return df


def _annual(df, how):
    df = df.copy()
    df['year'] = df['TIME_PERIOD'].astype(str).str[:4].astype(int)
    df = df.sort_values('TIME_PERIOD')
    g = df.groupby(['REF_AREA', 'year'])['OBS_VALUE']
    return (g.mean() if how == 'mean' else g.last()).unstack('REF_AREA')


def fetch_last_outturn_year(countries):
    """
    Last year with outturn data for the general government balance in Eurostat's EDP notification (gov_10dd_edpt1),
    the latest year across the given countries. The Commission forecast vintage year is this year plus one.
    """
    from .shocks import _eurostat
    geo = '+'.join(ISO2[c] if ISO2[c] != 'GR' else 'EL' for c in countries)
    df = _eurostat('gov_10dd_edpt1', f'A.PC_GDP.S13.B9.{geo}', 2015).dropna(subset=['OBS_VALUE'])
    return int(df.groupby('geo')['TIME_PERIOD'].max().max())


def fetch_benchmark_rates(countries, years):
    """
    Short-term (3M, GFS, last observation of year) and long-term (10Y, IRS, annual mean) benchmark rates.
    Countries without a short-term rate use the euro area rate. Returns dict code -> DataFrame (year x ISO3).
    """
    iso2 = {c: ISO2[c] for c in countries}
    inv = {v: k for k, v in iso2.items()}
    start = min(years)

    st_key = 'M.N.{}.W0.S13.S1.N.LI.LX.F3.S._Z.RT._T.F.V.A12._T'
    st = _annual(fetch_ecb('GFS', st_key.format('+'.join(list(iso2.values()) + ['I9'])) + f'?startPeriod={start}'), 'last')
    lt_frames = []
    for c in countries:
        cur = CURRENCY.get(c, 'EUR') if EURO_ADOPTION.get(c, 9999) > max(years) else 'EUR'
        try:
            lt_frames.append(_annual(fetch_ecb('IRS', f'M.{iso2[c]}.L.L40.CI.0000.{cur}.N.Z?startPeriod={start}'), 'mean'))
        except requests.HTTPError:
            lt_frames.append(_annual(fetch_ecb('IRS', f'M.{iso2[c]}.L.L40.CI.0000.EUR.N.Z?startPeriod={start}'), 'mean'))
    lt = pd.concat(lt_frames, axis=1)

    out = {}
    for code, df in [('INTEREST_RATE_ST', st), ('INTEREST_RATE_LT', lt)]:
        df = df.reindex(years)
        res = pd.DataFrame({inv[k]: df[k] for k in df.columns if k in inv}).reindex(columns=countries)
        if code == 'INTEREST_RATE_ST' and 'I9' in df.columns:
            for c in countries:
                if res[c].isna().all():
                    res[c] = df['I9']
        out[code] = res
    return out


def fetch_debt_structure(countries, start=2015):
    """
    Debt structure parameters from ECB GFS (annual):
        DEBT_ST_SHARE:              3-year average share of short-term debt (original maturity < 1 year)
        DEBT_LT_MATURING_SHARE:     latest share of long-term debt with residual maturity < 1 year
        DEBT_LT_MATURING_AVG_SHARE: 6-year average of the above
        DEBT_DOMESTIC_SHARE, DEBT_EUR_SHARE: latest currency shares
    Returns DataFrame (ISO3 x parameter) and the last year used.
    """
    iso2 = {c: ISO2[c] for c in countries}
    keys = {
        'TOTAL': 'A.N.{}.W0.S13.S1.C.L.LE.GD.T._Z.XDC._T.F.V.N._T',
        'ST': 'A.N.{}.W0.S13.S1.C.L.LE.GD.S._Z.XDC._T.F.V.N._T',
        'LT': 'A.N.{}.W0.S13.S1.C.L.LE.GD.L._Z.XDC._T.F.V.N._T',
        'MATURING': 'A.N.{}.W0.S13.S1.C.L.LE.GD.TS._Z.XDC._T.F.V.N._T',
        'DOMESTIC': 'A.N.{}.W0.S13.S1.C.L.LE.GD.T._Z.XDC.XDC.F.V.N._T',
        'EUR': 'A.N.{}.W0.S13.S1.C.L.LE.GD.T._Z.XDC.XPC.F.V.N._T',
    }
    area = '+'.join(iso2.values())
    d = {k: _annual(fetch_ecb('GFS', key.format(area) + f'?startPeriod={start}'), 'last') for k, key in keys.items()}
    rows, last_years = {}, {}
    for c, a in iso2.items():
        get = lambda k: d[k][a] if a in d[k].columns else pd.Series(dtype=float)
        total, st, lt, mat = get('TOTAL'), get('ST'), get('LT'), get('MATURING')
        last = total.last_valid_index()
        st_share = (st / total).dropna()
        lt_mat = ((mat - st) / lt).dropna()
        dom, eur = (get('DOMESTIC') / total).dropna(), (get('EUR') / total).dropna()
        rows[c] = {
            'DEBT_ST_SHARE': st_share.loc[last - 2:last].mean() if len(st_share) else np.nan,
            'DEBT_LT_MATURING_SHARE': lt_mat.iloc[-1] if len(lt_mat) else np.nan,
            'DEBT_LT_MATURING_AVG_SHARE': lt_mat.loc[last - 5:last].mean() if len(lt_mat) else np.nan,
            'DEBT_DOMESTIC_SHARE': dom.iloc[-1] if len(dom) else np.nan,
            'DEBT_EUR_SHARE': eur.iloc[-1] if len(eur) else 0.0,
        }
        last_years[c] = last
    return pd.DataFrame(rows).T, last_years


# ========================================================================================= #
#                                            ESM                                            #
# ========================================================================================= #

def fetch_esm_repayments(countries, start_year):
    """
    ESM/EFSF repayment schedules (latest events) in bn EUR by year from start_year.
    Returns DataFrame (year x ISO3), zero where no repayments.
    """
    df = pd.read_csv(io.StringIO(_get(ESM_URL).text))
    df = df.loc[df['Event'].astype(str).str.startswith('Latest')]
    df['year'] = pd.to_datetime(df['Payment Date'], dayfirst=True).dt.year
    df['value'] = pd.to_numeric(df['Payment Amount'], errors='coerce') / 1e9
    df = df.loc[df['year'] >= start_year]
    inv = {v: k for k, v in ISO2.items()}
    df['country'] = df['Country'].map(inv)
    out = df.groupby(['year', 'country'])['value'].sum().unstack('country')
    return out.reindex(columns=countries).fillna(0.0)


# ========================================================================================= #
#                                    Reference files                                        #
# ========================================================================================= #

def read_awg_table(path, title):
    """
    Read one cross-country table of the Ageing Report statistical annex. Returns DataFrame (year x ISO3).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['cross-country_tables']
    rows = ws.iter_rows(values_only=True)
    for r in rows:
        if any(isinstance(v, str) and v.strip().startswith(title) for v in r):
            break
    else:
        raise KeyError(f"Table '{title}' not found in {path}")
    header, data = None, {}
    inv = {v: k for k, v in EC_ISO2.items()}
    for r in rows:
        vals = [v for v in r if v is not None]
        if header is None:
            if any(isinstance(v, (int, float)) and 2000 < v < 2100 for v in vals):
                j0 = next(j for j, v in enumerate(r) if isinstance(v, (int, float)) and 2000 < v < 2100)
                header = {j: int(v) for j, v in enumerate(r) if j >= j0 and isinstance(v, (int, float))}
            continue
        label = next((v for v in r if isinstance(v, str)), None)
        if label is None or label.strip() in ('EA', 'EU'):
            break
        if label.strip() in inv:
            data[inv[label.strip()]] = {y: r[j] for j, y in header.items() if isinstance(r[j], (int, float))}
    wb.close()
    return pd.DataFrame(data)


def read_dsm_fiche_rows(path, countries, patterns):
    """
    Read rows of the DSM country fiche baseline table by label pattern.
    Returns {code: DataFrame (year x ISO3)}.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {code: {} for code in patterns}
    for c in countries:
        ws = wb[EC_ISO2[c]]
        header = None
        for r in ws.iter_rows(values_only=True):
            label = next((str(v) for v in r if isinstance(v, str)), '')
            if header is None and re.search(r'baseline scenario', label):
                years = {j: int(v) for j, v in enumerate(r) if isinstance(v, (str, int)) and str(v).isdigit()}
                header = years if len(years) >= 5 else None
                continue
            if header is None:
                continue
            for code, pat in patterns.items():
                if c not in out[code] and any(isinstance(v, str) and re.search(pat, v) for v in r):
                    out[code][c] = {y: pd.to_numeric(r[j], errors='coerce') for j, y in header.items()}
    wb.close()
    return {code: pd.DataFrame(v) for code, v in out.items()}


def read_ogwg(path, countries):
    """
    Read a user-supplied Output Gaps Working Group file with sheets 'Real GDP' and 'Pot_GDP'
    (years in rows, Commission country codes in columns, as distributed by the OGWG).
    Returns {'REAL_GDP': DataFrame, 'POTENTIAL_GDP': DataFrame} (year x ISO3).
    """
    out = {}
    for code, sheet in [('REAL_GDP', 'Real GDP'), ('POTENTIAL_GDP', 'Pot_GDP')]:
        df = pd.read_excel(path, sheet_name=sheet, skiprows=1, index_col=0)
        df = df.loc[pd.to_numeric(df.index, errors='coerce').notna()]
        df.index = df.index.astype(int)
        out[code] = pd.DataFrame({c: pd.to_numeric(df[EC_ISO2[c]], errors='coerce') for c in countries
                                  if EC_ISO2[c] in df.columns})
    return out
