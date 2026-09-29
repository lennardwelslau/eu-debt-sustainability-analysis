# ========================================================================================= #
#               European Commission Debt Sustainability Analysis - Results Tables           #
# ========================================================================================= #
#
# Formats model results as labelled pandas tables (for display in notebooks) and writes them
# to a single, readable Excel workbook (for GroupDsaModel.save_results).
#
# Author: Lennard Welslau
# ========================================================================================= #

import numpy as np
import pandas as pd

# Key model variables and labels
KEY_VARIABLES = {
    'd': 'Debt (% of GDP)',
    'ob': 'Overall balance (% of GDP)',
    'pb': 'Primary balance (% of GDP)',
    'sb': 'Structural balance (% of GDP)',
    'spb': 'Structural primary balance (% of GDP)',
    'spb_bca': 'SPB before change in ageing costs (% of GDP)',
    'spb_bca_adjustment': 'Annual SPB adjustment (pp.)',
    'net_expenditure_growth': 'Net expenditure growth (%)',
    'ageing_component': 'Change in ageing costs after adjustment (pp.)',
    'cyclical_component': 'Cyclical component (% of GDP)',
    'interest_ratio': 'Interest expenditure (% of GDP)',
    'iir': 'Implicit interest rate (%)',
    'i_st': 'Short-term market interest rate (%)',
    'i_lt': 'Long-term market interest rate (%)',
    'rg': 'Real GDP growth (%)',
    'rg_pot': 'Potential GDP growth (%)',
    'output_gap': 'Output gap (% of potential GDP)',
    'pi': 'GDP deflator inflation (%)',
    'ng': 'Nominal GDP growth (%)',
    'sf': 'Stock-flow adjustment (% of GDP)',
    'gfn': 'Gross financing needs (% of GDP)',
}

SCENARIO_LABELS = {
    'main_adjustment': 'Adjustment scenario',
    'lower_spb': 'Lower SPB scenario',
    'financial_stress': 'Financial stress scenario',
    'adverse_r_g': 'Adverse r-g scenario',
}

COUNTRY_NAMES = {
    'AUT': 'Austria', 'BEL': 'Belgium', 'BGR': 'Bulgaria', 'HRV': 'Croatia', 'CYP': 'Cyprus', 'CZE': 'Czechia',
    'DNK': 'Denmark', 'EST': 'Estonia', 'FIN': 'Finland', 'FRA': 'France', 'DEU': 'Germany', 'GRC': 'Greece',
    'HUN': 'Hungary', 'IRL': 'Ireland', 'ITA': 'Italy', 'LVA': 'Latvia', 'LTU': 'Lithuania', 'LUX': 'Luxembourg',
    'MLT': 'Malta', 'NLD': 'Netherlands', 'POL': 'Poland', 'PRT': 'Portugal', 'ROU': 'Romania', 'SVK': 'Slovakia',
    'SVN': 'Slovenia', 'ESP': 'Spain', 'SWE': 'Sweden',
}


def criterion_label(key):
    """
    Readable label for an SPB target key.
    """
    fixed = {
        'deficit_reduction': 'Deficit below 3% (DSA)',
        'stochastic': 'Stochastic (DSA)',
        'edp': 'Excessive deficit procedure',
        'debt_safeguard': 'Debt sustainability safeguard',
        'deficit_resilience': 'Deficit resilience safeguard',
        'binding': 'Binding',
        'floor': 'Floor on annual adjustment',
    }
    if key in fixed:
        return fixed[key]
    if key in SCENARIO_LABELS:
        return f'{SCENARIO_LABELS[key]}: debt declines or below 60% (DSA)'
    for prefix, text in [('debt_declines_', 'debt declines'), ('debt_below_60_', 'debt below 60%')]:
        if key.startswith(prefix) and key[len(prefix):] in SCENARIO_LABELS:
            return f'{SCENARIO_LABELS[key[len(prefix):]]}: {text} (DSA)'
    return key


def key_variables(df, variables=None):
    """
    Labelled table of key variables (index: year) from a model DataFrame as returned by model.df(all=True).
    """
    df = df.reset_index().set_index('y')
    df.index.name = 'Year'
    if 'GFN' in df.columns and 'ngdp' in df.columns:
        df['gfn'] = df['GFN'] / df['ngdp'] * 100
    variables = variables or list(KEY_VARIABLES)
    cols = [v for v in variables if v in df.columns]
    return df[cols].rename(columns=KEY_VARIABLES)


def display_tables(tables):
    """
    Display tables as HTML in notebooks, as text otherwise.
    """
    try:
        from IPython import get_ipython
        from IPython.display import display, Markdown
        if get_ipython() is not None:
            for title, table in tables.items():
                display(Markdown(f'**{title}**'))
                display(table.style.format(precision=2, na_rep='') if isinstance(table, pd.DataFrame) else table)
            return
    except ImportError:
        pass
    with pd.option_context('display.float_format', '{:,.2f}'.format, 'display.width', 200, 'display.max_columns', 30):
        for title, table in tables.items():
            print(f'\n{title}\n{"-" * len(title)}')
            if isinstance(table, pd.DataFrame):
                table = table.map(lambda x: f'{x:,.2f}' if isinstance(x, float) and not np.isnan(x)
                                  else '' if isinstance(x, float) else x)
            print(table.to_string())


# ========================================================================================= #
#                                     EXCEL OUTPUT                                          #
# ========================================================================================= #

def write_results_workbook(path, results, meta):
    """
    Write group results to one Excel workbook.

    Parameters:
        path: output path
        results: {country: {'summary': Series, 'targets': DataFrame, 'path': DataFrame,
                            'scenarios': {scenario: key variable DataFrame}}}
        meta: dict with description entries for the README sheet
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    countries = [c for c in results if results[c]]
    summary = pd.concat({c: results[c]['summary'] for c in countries if 'summary' in results[c]}, axis=1).T
    targets = pd.concat({c: results[c]['targets']['SPB at end of adjustment'] for c in countries
                         if 'targets' in results[c]}, axis=1).T
    paths = pd.concat({c: results[c]['path'] for c in countries if 'path' in results[c]}, names=['Country'])

    debt = {}
    for c in countries:
        for scenario, df in results[c].get('scenarios', {}).items():
            debt.setdefault(scenario, {})[c] = df[KEY_VARIABLES['d']]

    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        readme = pd.DataFrame(list(meta.items()) + [
            ('', ''),
            ('Summary', 'Binding SPB target, annual adjustment and binding criterion by country'),
            ('SPB targets', 'SPB at the end of the adjustment period required by each criterion (% of GDP)'),
            ('Adjustment paths', 'Annual SPB adjustment, minimum steps, balances, debt and net expenditure growth'),
            ('Debt <scenario>', 'Debt ratio by year and country'),
            ('<ISO3>', 'Key variables by year for the binding and no-policy-change scenarios'),
        ], columns=['Item', 'Description'])
        readme.to_excel(writer, sheet_name='README', index=False)
        if len(summary):
            summary.index.name = 'ISO'
            summary.to_excel(writer, sheet_name='Summary')
        if len(targets):
            targets.index.name = 'ISO'
            targets.to_excel(writer, sheet_name='SPB targets')
        if len(paths):
            paths.to_excel(writer, sheet_name='Adjustment paths', merge_cells=False)
        for scenario, d in debt.items():
            pd.DataFrame(d).to_excel(writer, sheet_name=f'Debt {scenario}'[:31])
        for c in countries:
            scenarios = results[c].get('scenarios', {})
            if scenarios:
                table = pd.concat({s: df.T for s, df in scenarios.items()}, names=['Scenario', 'Variable'])
                table.to_excel(writer, sheet_name=c, merge_cells=False)

        # Formatting
        header_fill, header_font = PatternFill('solid', fgColor='1F4E78'), Font(bold=True, color='FFFFFF')
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal='center', wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = '0.00'
            for j in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(j)].width = 12
            ws.column_dimensions['A'].width = 16 if ws.title != 'README' else 20
            if ws.title in COUNTRY_NAMES:
                ws.column_dimensions['B'].width = 50
                ws.freeze_panes = 'C2'
            elif ws.title == 'README':
                ws.column_dimensions['B'].width = 100
            elif ws.title == 'Adjustment paths':
                ws.freeze_panes = 'C2'
            else:
                ws.freeze_panes = 'B2'
    return path
